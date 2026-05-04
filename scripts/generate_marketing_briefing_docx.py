# One-off / reusable: generates Marketing briefing Word doc for IHSG Supertrend Scanner.
# Run: python scripts/generate_marketing_briefing_docx.py

from pathlib import Path

from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "IHSG_Supertrend_Scanner_Briefing_Marketing.docx"


def add_body_style(doc: Document):
    style = doc.styles["Normal"]
    font = style.font
    font.name = "Calibri"
    font.size = Pt(11)


def main():
    doc = Document()
    add_body_style(doc)

    t = doc.add_heading("IHSG Supertrend Scanner", 0)
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph("Briefing untuk tim marketing & storytelling (orang awam)").alignment = (
        WD_ALIGN_PARAGRAPH.CENTER
    )
    doc.add_paragraph("")

    doc.add_heading("1. Elevator pitch (satu kalimat)", level=1)
    doc.add_paragraph(
        "Bot ini memantau ratusan saham IHSG secara otomatis, menghitung pola teknikal "
        "(utamanya Supertrend + skor gabungan), lalu mengirim pemberitahuan ke Telegram "
        "hanya saat ada sinyal baru yang lolos filter likuiditas—bukan robot yang membeli/jual "
        "otomatis di sekuritas."
    )

    doc.add_heading("2. Cerita singkat (storytelling)", level=1)
    doc.add_paragraph(
        "Bayangkan ada asisten riset yang setiap hari kerja membuka chart ratusan saham sekaligus "
        "dan mencari pola yang dipakai trader teknikal: apakah tren lagi “hijau”, apakah volume "
        "mencurigakan (sering dikaitkan dengan akumulasi), apakah ada koreksi yang masih “sehat”, "
        "atau ada petunjuk divergensi (harga turun tapi momentum mulai membaik)."
    )
    doc.add_paragraph(
        "Bot ini membaca data harga harian, menghitung indikator, memberi nilai (skor), dan "
        "mengelompokkan ke dalam beberapa jenis sinyal. Jika ada yang baru dan belum di-alert "
        "hari itu, pengguna mendapat notifikasi di Telegram—tanpa harus manual refresh ratusan ticker."
    )
    doc.add_paragraph(
        "Analogi aman untuk marketing: ini seperti alarm pintar untuk watchlist besar, bukan "
        "pengemudi mobil yang mengambil alih setir."
    )

    doc.add_heading("3. Alur kerja sistem (ringkas)", level=1)
    steps = [
        (
            "Ambil data",
            "Harga dari Yahoo Finance (yfinance), periode ~90 hari, interval 1 hari (daily) untuk semua sinyal.",
        ),
        (
            "Filter likuiditas",
            "Hanya saham dengan rata-rata nilai transaksi 5 hari ≥ 5 miliar Rupiah yang masuk daftar sinyal "
            "(mengurangi noise saham terlalu illiquid).",
        ),
        (
            "Indikator & Supertrend",
            "Supertrend mengikuti logika yang diselaraskan dengan Pine Script v3 (ATR, pita, arah bullish/bearish, "
            "deteksi breakout). Indikator lain: EMA, volume, Stochastic RSI, ADX, MACD, OBV, pola candle, "
            "support/resistance, ATR untuk target, dll.",
        ),
        (
            "Skor total (versi v5)",
            "Skor 0–100 dari blok: tren, regime pasar (trending/volatile), volume, momentum, posisi vs EMA, pola—"
            "dengan penalti untuk sideways dan bearish divergence saat tren masih bullish. Label contoh: "
            "STRONG BUY, ACCUMULATE, HOLD, AVOID.",
        ),
        (
            "Empat jenis sinyal ke Telegram",
            "Strong Buy (breakout terkonfirmasi + skor tinggi); Accumulation (Stoch + volume tidak biasa dalam tren "
            "bullish); Early entry / “serok bawah” (koreksi terkendali dalam uptrend); Bullish divergence "
            "(potensi pembalikan, dengan grade kekuatan).",
        ),
        (
            "Target di pesan",
            "TP1 dan TP2 dari ATR dan/atau resistance (dengan aturan agar TP2 realistis). SL ditampilkan di pesan "
            "(contoh −5% dari harga terkini sebagai referensi tampilan; bukan eksekusi otomatis).",
        ),
        (
            "Anti-spam",
            "Status per saham disimpan; sinyal yang sudah di-alert hari itu tidak diulang berlebihan.",
        ),
        (
            "Jadwal",
            "Scheduler “smart sleep”: aktif jam pasar, scan rutin ~1 menit; recap pembukaan/penutupan; evaluasi "
            "hasil sinyal sekitar 16:30 WIB (rekapan TP/SL intraday untuk sinyal yang dilacak).",
        ),
        (
            "Learning (opsional)",
            "Jika PostgreSQL tersedia, sinyal dapat dicatat dan regime pasar dipakai untuk pelacakan; tanpa DB, "
            "bot tetap berjalan.",
        ),
    ]
    for title, text in steps:
        p = doc.add_paragraph()
        p.add_run(f"{title}. ").bold = True
        p.add_run(text)

    doc.add_heading("4. Glosarium: bahasa marketing → arti teknis", level=1)
    table = doc.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    hdr[0].text = "Yang dipromosikan"
    hdr[1].text = "Arti di produk"
    for cell in hdr:
        for p in cell.paragraphs:
            for r in p.runs:
                r.bold = True
    rows = [
        (
            "“AI trading”",
            "Lebih tepat: aturan teknikal + skor berbobot di kode; bukan model ML yang menebak sentimen berita.",
        ),
        ("“Scan ratusan saham”", "Benar: daftar besar ticker IHSG, data di-batch lewat Yahoo."),
        (
            "“Sinyal real-time”",
            "Notifikasi bisa sering intraday; sinyal berbasis candle harian dari sumber data.",
        ),
        (
            "“Anti noise”",
            "Filter turnover, konfirmasi breakout, penalti sideways, syarat volume pada accumulation, dll.",
        ),
        ("“Bukan sekuritas”", "Tidak terhubung ke RDN; tidak mengeksekusi order beli/jual."),
    ]
    for a, b in rows:
        row = table.add_row().cells
        row[0].text = a
        row[1].text = b

    doc.add_heading("5. Posisi produk (monetisasi)", level=1)
    doc.add_paragraph(
        "Nilai yang dijual: hemat waktu, konsistensi aturan, cakupan luas IHSG, notifikasi terstruktur "
        "(skor + TP/SL), disiplin “hanya alert yang baru”."
    )
    doc.add_paragraph(
        "Komunikasi yang aman: alat bantu analisa/alert; keputusan investasi tetap pengguna; performa bergantung "
        "pada kondisi pasar, parameter, dan data pihak ketiga; hindari klaim pasti profit atau tanpa risiko."
    )

    doc.add_heading("6. Checklist tim marketing sebelum copywriting", level=1)
    for item in [
        "Tekankan scanner + Telegram alert, bukan auto-trade.",
        "Sebut cakupan IHSG dan filter likuiditas.",
        "Jelaskan 4 jenis sinyal dengan bahasa sederhana.",
        "Sertakan disclaimer investasi sesuai praktik promosi di Indonesia jika dijual ke publik.",
    ]:
        doc.add_paragraph(item, style="List Bullet")

    doc.add_heading("7. Catatan keamanan (internal)", level=1)
    doc.add_paragraph(
        "Kredensial Telegram sebaiknya via environment variable di produksi. Jangan menyebar token/chat ID di "
        "materi publik; rotasi token jika pernah terpapar."
    )

    doc.add_paragraph("")
    p = doc.add_paragraph(f"Dokumen dihasilkan otomatis. File: {OUT.name}")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.save(OUT)
    print(f"Written: {OUT}")


if __name__ == "__main__":
    main()
